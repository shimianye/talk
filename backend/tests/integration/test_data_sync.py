"""种子 upsert、manifest 快速路径和知识增量更新集成测试。"""
from __future__ import annotations

import asyncio
import os
import shutil
from pathlib import Path

import openpyxl
import pytest
from alembic import command
from alembic.config import Config
from sqlalchemy import delete, func, select
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.rag.embedding import MockEmbeddingProvider
from app.models import KnowledgeChunk, KnowledgeDocument, Product
from app.services.kb_sync import sync_knowledge_docs
from app.services.seed_sync import DEFAULT_DATA_DIR, sync_seed_data

BACKEND_DIR = Path(__file__).resolve().parents[2]
KB_DIR = BACKEND_DIR.parent / "data" / "kb-docs"


class CountingEmbedding(MockEmbeddingProvider):
    """记录实际参与向量化的文本数，验证快速路径没有重复调用。"""

    def __init__(self, dim: int = 1024):
        super().__init__(dim)
        self.text_count = 0

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.text_count += len(texts)
        return await super().embed(texts)


class BrokenEmbedding(MockEmbeddingProvider):
    """故意返回错误数量，验证调用方事务会回滚半成品。"""

    async def embed(self, texts: list[str]) -> list[list[float]]:
        vectors = await super().embed(texts)
        return vectors[:-1]


def _test_url() -> str:
    """只允许显式迁移测试库，避免 reset 误伤主库。"""
    url = os.getenv("ALEMBIC_DATABASE_URL")
    if not url:
        pytest.skip("未设置 ALEMBIC_DATABASE_URL，跳过数据同步集成测试")
    if not (make_url(url).database or "").endswith("_migration_test"):
        pytest.fail("数据同步集成测试只允许使用 _migration_test 数据库")
    return url


def _copy_seed_data(target: Path) -> Path:
    """复制种子工作簿，允许测试修改而不触碰仓库源文件。"""
    target.mkdir()
    for source in DEFAULT_DATA_DIR.glob("*.xlsx"):
        shutil.copy2(source, target / source.name)
    return target


def _copy_two_docs(target: Path) -> Path:
    """复制两篇知识文档，控制集成测试向量化成本。"""
    (target / "apple").mkdir(parents=True)
    for filename in ("iphone-16.md", "iphone-15.md"):
        shutil.copy2(KB_DIR / "apple" / filename, target / "apple" / filename)
    return target


async def test_seed_and_knowledge_sync_are_repeatable(tmp_path: Path):
    """验证首次导入、快速跳过、内容更新、孤儿清理和显式 reset。"""
    url = _test_url()
    config = Config(str(BACKEND_DIR / "alembic.ini"))
    await asyncio.to_thread(command.downgrade, config, "base")
    await asyncio.to_thread(command.upgrade, config, "head")

    seed_dir = _copy_seed_data(tmp_path / "seed")
    docs_dir = _copy_two_docs(tmp_path / "docs")
    embedding = CountingEmbedding()
    engine = create_async_engine(url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        async with factory() as session, session.begin():
            first_seed = await sync_seed_data(session, seed_dir)
            first_kb = await sync_knowledge_docs(session, docs_dir, embedding)
        first_embedding_count = embedding.text_count
        assert first_seed["skipped"] is False
        assert first_kb["documents"] == 2
        assert first_embedding_count > 0

        async with factory() as session, session.begin():
            second_seed = await sync_seed_data(session, seed_dir)
            second_kb = await sync_knowledge_docs(session, docs_dir, embedding)
        assert second_seed["skipped"] is True
        assert second_kb == {"documents": 0, "chunks": 0, "skipped": 2, "deleted": 0}
        assert embedding.text_count == first_embedding_count

        async with factory() as session, session.begin():
            document_id = await session.scalar(
                select(KnowledgeDocument.document_id).limit(1)
            )
            await session.execute(
                delete(KnowledgeChunk).where(
                    KnowledgeChunk.document_id == document_id
                )
            )
        before_repair = embedding.text_count
        async with factory() as session, session.begin():
            repaired_kb = await sync_knowledge_docs(session, docs_dir, embedding)
        assert repaired_kb["documents"] == 1
        assert repaired_kb["skipped"] == 1
        assert embedding.text_count > before_repair

        async with factory() as session, session.begin():
            missing_product_id = await session.scalar(select(Product.product_id).limit(1))
            await session.execute(
                delete(Product).where(Product.product_id == missing_product_id)
            )
        async with factory() as session, session.begin():
            repaired_seed = await sync_seed_data(session, seed_dir)
        assert repaired_seed["skipped"] is False
        assert repaired_seed["database_counts"]["products"] == 9

        workbook_path = seed_dir / "phone_specs.xlsx"
        workbook = openpyxl.load_workbook(workbook_path)
        worksheet = workbook["products"]
        headers = [cell.value for cell in worksheet[1]]
        brand_column = headers.index("brand") + 1
        product_id_column = headers.index("product_id") + 1
        product_id = str(worksheet.cell(2, product_id_column).value)
        worksheet.cell(2, brand_column).value = "Apple-Test"
        workbook.save(workbook_path)
        workbook.close()

        async with factory() as session, session.begin():
            changed_seed = await sync_seed_data(session, seed_dir)
            brand = await session.scalar(
                select(Product.brand).where(Product.product_id == product_id)
            )
        assert changed_seed["skipped"] is False
        assert brand == "Apple-Test"

        changed_doc = docs_dir / "apple" / "iphone-16.md"
        changed_doc.write_text(
            changed_doc.read_text(encoding="utf-8") + "\n\n测试版本变化。",
            encoding="utf-8",
        )
        async with factory() as session, session.begin():
            changed_kb = await sync_knowledge_docs(
                session, docs_dir, embedding, chunk_size=400
            )
        assert changed_kb["documents"] == 2
        assert changed_kb["skipped"] == 0

        (docs_dir / "apple" / "iphone-15.md").unlink()
        async with factory() as session, session.begin():
            deleted_kb = await sync_knowledge_docs(
                session, docs_dir, embedding, chunk_size=400
            )
            managed_count = int(
                await session.scalar(
                    select(func.count()).select_from(KnowledgeDocument)
                ) or 0
            )
        assert deleted_kb["deleted"] == 1
        assert managed_count == 1

        broken_doc = docs_dir / "apple" / "broken.md"
        broken_doc.write_text("# 事务失败测试\n\n此文档不应留下半成品。", encoding="utf-8")
        with pytest.raises(ValueError, match="Embedding 返回数量"):
            async with factory() as session, session.begin():
                await sync_knowledge_docs(session, docs_dir, BrokenEmbedding())
        async with factory() as session:
            broken_count = int(
                await session.scalar(
                    select(func.count()).select_from(KnowledgeDocument).where(
                        KnowledgeDocument.file_path == "apple/broken.md"
                    )
                ) or 0
            )
        assert broken_count == 0

        async with factory() as session, session.begin():
            reset_seed = await sync_seed_data(session, seed_dir, reset=True)
        assert reset_seed["skipped"] is False
        assert reset_seed["database_counts"]["products"] == 9
    finally:
        await engine.dispose()
