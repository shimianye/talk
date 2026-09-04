# -*- coding: utf-8 -*-
"""为 evaluation_dataset.xlsx 增加 expected_tool_name 和 expected_route 两列"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

BASE = r"C:\Users\Administrator\Doubao\chats\2026-09-03\new-chat-5"

# intent -> (expected_tool_name, expected_route)
intent_map = {
    "product_spec":      ("无", "RAG"),
    "product_compare":   ("无", "RAG"),
    "price_query":       ("query_price", "API"),
    "inventory_query":   ("query_inventory", "API"),
    "order_query":       ("query_order", "API"),
    "refund_policy":     ("无", "RAG"),
    "warranty_policy":   ("无", "RAG"),
    "shipping_policy":   ("无", "RAG"),
    "promotion_query":   ("query_promotion", "API"),
    "purchase_recommend":("无", "RAG"),
    "invoice_query":     ("无", "RAG"),
    "complaint":         ("无", "人工"),
    "after_sales_apply": ("create_after_sales", "API"),
    "account_issue":     ("无", "RAG"),
    "unknown":           ("无", "通用回复"),
    "greeting":          ("无", "通用回复"),
    "store_info":        ("无", "RAG"),
    "escalation":        ("无", "人工"),
}

# 特殊覆盖：某些问题虽然intent是order_query但should_transfer_human=是，路由应为人工
# 某些越权问题路由为"拒答"

wb = openpyxl.load_workbook(os.path.join(BASE, "evaluation_dataset.xlsx"))
ws = wb["evaluation"]

# 读取现有表头
headers = [c.value for c in ws[1]]
print("现有列:", headers)

# 找到 intent 列和 should_transfer_human 列
intent_col = headers.index("intent") + 1
transfer_col = headers.index("should_transfer_human") + 1
question_col = headers.index("question") + 1

# 添加两列表头
new_col1 = ws.max_column + 1
new_col2 = ws.max_column + 2
ws.cell(1, new_col1, "expected_tool_name")
ws.cell(1, new_col2, "expected_route")

# 样式
hf = Font(bold=True, color="FFFFFF", size=11)
hfill = PatternFill("solid", fgColor="2F5496")
ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
bd = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
for c in [new_col1, new_col2]:
    cell = ws.cell(1, c)
    cell.font = hf; cell.fill = hfill; cell.alignment = ha; cell.border = bd

# 越权关键词
privacy_keywords = ["别人的订单", "他人", "查一下别人", "老板电话", "内部价", "律师函", "爆炸", "人身安全"]

count = 0
for row in range(2, ws.max_row + 1):
    intent = ws.cell(row, intent_col).value
    transfer = ws.cell(row, transfer_col).value
    question = ws.cell(row, question_col).value or ""

    tool, route = intent_map.get(intent, ("无", "RAG"))

    # 越权/隐私问题 -> 拒答
    if any(kw in question for kw in ["别人的订单", "他人的订单", "查一下别人", "查别人"]):
        tool = "无"
        route = "拒答"
    # 投诉/法律威胁/安全事故 -> 人工
    elif any(kw in question for kw in ["律师函", "告你们", "爆炸", "赔偿", "投诉", "态度差"]):
        if intent in ["complaint", "escalation"] or transfer == "是":
            tool = "无"
            route = "人工"
    # should_transfer_human=是 -> 人工
    elif transfer == "是":
        tool = "无"
        route = "人工"

    ws.cell(row, new_col1, tool)
    ws.cell(row, new_col2, route)
    ws.cell(row, new_col1).border = bd
    ws.cell(row, new_col2).border = bd
    ws.cell(row, new_col1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.cell(row, new_col2).alignment = Alignment(vertical="center", wrap_text=True)
    count += 1

# 设置列宽
ws.column_dimensions[openpyxl.utils.get_column_letter(new_col1)].width = 22
ws.column_dimensions[openpyxl.utils.get_column_letter(new_col2)].width = 14

wb.save(os.path.join(BASE, "evaluation_dataset.xlsx"))
print(f"已为 {count} 条评测题补充 expected_tool_name 和 expected_route")

# 统计路由分布
wb2 = openpyxl.load_workbook(os.path.join(BASE, "evaluation_dataset.xlsx"))
ws2 = wb2["evaluation"]
routes = {}
tools = {}
for row in range(2, ws2.max_row + 1):
    r = ws2.cell(row, new_col2).value
    t = ws2.cell(row, new_col1).value
    routes[r] = routes.get(r, 0) + 1
    tools[t] = tools.get(t, 0) + 1
print(f"路由分布: {routes}")
print(f"工具分布: {tools}")
