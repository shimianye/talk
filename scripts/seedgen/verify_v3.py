# -*- coding: utf-8 -*-
import openpyxl, os
from datetime import datetime

BASE = r"C:\Users\Administrator\Doubao\chats\2026-09-03\new-chat-5"
os.chdir(BASE)

wb = openpyxl.load_workbook('business_data.xlsx')
print('=== business_data.xlsx ===')
for s in wb.sheetnames:
    ws = wb[s]
    h = [c.value for c in ws[1]]
    extra = ''
    if s == 'store_products':
        extra = ' | price_type=%s original_price=%s promotion_price=%s' % (
            'price_type' in h, 'original_price' in h, 'promotion_price' in h)
    if s == 'logistics':
        extra = ' | package_id=%s is_split_shipment=%s' % (
            'package_id' in h, 'is_split_shipment' in h)
    print('  %s: %d rows, %d cols%s' % (s, ws.max_row-1, ws.max_column, extra))

ws = wb['orders']
errors = 0
def parse(s):
    return datetime.strptime(s, '%Y-%m-%d %H:%M:%S') if s else None
for row in ws.iter_rows(min_row=2, values_only=True):
    ct, pt, st, rt = row[12], row[13], row[14], row[15]
    c, p, sh, r = parse(ct), parse(pt), parse(st), parse(rt)
    if p and c and p < c: errors += 1
    if sh and p and sh < p: errors += 1
    if r and sh and r < sh: errors += 1
print('  订单时间逻辑错误: %d' % errors)

ws6 = wb['after_sales_cases']
mismatch = 0
for row in ws6.iter_rows(min_row=2, values_only=True):
    ct, st, desc = row[3], row[4], row[5]
    kw_map = {'退货':'退', '换货':'换', '维修':'修', '退款':'退款',
              '投诉':'投诉', '碎屏保障':'碎屏', '价保':'降价'}
    if ct in kw_map:
        if ct == '价保':
            if '降价' not in desc and '差价' not in desc and '价保' not in desc:
                mismatch += 1
        elif kw_map[ct] not in desc:
            mismatch += 1
print('  售后分类不一致: %d' % mismatch)

wb2 = openpyxl.load_workbook('evaluation_dataset.xlsx')
ws2 = wb2['evaluation']
h2 = [c.value for c in ws2[1]]
print('\n=== evaluation_dataset.xlsx ===')
print('  列数: %d, expected_tool_name=%s, expected_route=%s' % (
    len(h2), 'expected_tool_name' in h2, 'expected_route' in h2))
print('  数据行: %d' % (ws2.max_row-1))

print('\n=== kb-docs markdown ===')
total = 0
for root, dirs, files in os.walk('kb-docs'):
    for f in sorted(files):
        if f.endswith('.md'):
            total += 1
            print('  %s' % os.path.join(root, f))
print('  总计: %d 份' % total)
